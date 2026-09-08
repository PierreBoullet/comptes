#!/usr/bin/env python3
"""Export every worksheet from Financements.xlsx to CSV files."""

from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = "".join(char for char in normalized if not unicodedata.combining(char))
    slug = re.sub(r"[^A-Za-z0-9]+", "_", ascii_value).strip("_").lower()
    return slug or "sheet"


def cell_content(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    return value


def export_sheet_grids(workbook_path: Path, output_dir: Path) -> list[Path]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    sheet_dir = output_dir / "onglets"
    sheet_dir.mkdir(parents=True, exist_ok=True)

    created: list[Path] = []
    for worksheet in workbook.worksheets:
        output_path = sheet_dir / f"{slugify(worksheet.title)}.csv"
        with output_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
            writer = csv.writer(csv_file, delimiter=";")
            for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, max_col=worksheet.max_column):
                writer.writerow([cell_content(cell.value) for cell in row])
        created.append(output_path)
    return created


def export_all_cells(workbook_path: Path, output_dir: Path) -> Path:
    workbook_formulas = load_workbook(workbook_path, read_only=True, data_only=False)
    workbook_values = load_workbook(workbook_path, read_only=True, data_only=True)

    output_path = output_dir / "Financements_all_cells.csv"
    with output_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.writer(csv_file, delimiter=";")
        writer.writerow(["Feuille", "Ligne", "Colonne", "Adresse", "Valeur", "Formule"])

        for worksheet_formulas in workbook_formulas.worksheets:
            worksheet_values = workbook_values[worksheet_formulas.title]
            for row_index, row in enumerate(worksheet_formulas.iter_rows(
                min_row=1,
                max_row=worksheet_formulas.max_row,
                max_col=worksheet_formulas.max_column,
            ), start=1):
                for column_index, formula_cell in enumerate(row, start=1):
                    coordinate = f"{get_column_letter(column_index)}{row_index}"
                    raw_formula_value = formula_cell.value
                    value_cell = worksheet_values[coordinate]
                    formula = raw_formula_value if isinstance(raw_formula_value, str) and raw_formula_value.startswith("=") else None
                    value = value_cell.value if formula else formula_cell.value
                    if value is None and formula is None:
                        continue
                    writer.writerow([
                        worksheet_formulas.title,
                        row_index,
                        get_column_letter(column_index),
                        coordinate,
                        cell_content(value),
                        formula,
                    ])
    return output_path


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Export all Excel workbook sheets to CSV")
    parser.add_argument("--workbook", default="Financements.xlsx", help="Excel workbook to export")
    parser.add_argument("--output-dir", default="exports", help="Directory where CSV files are written")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    workbook_path = Path(args.workbook)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    all_cells_path = export_all_cells(workbook_path, output_dir)
    sheet_paths = export_sheet_grids(workbook_path, output_dir)

    print(f"Global CSV: {all_cells_path}")
    for sheet_path in sheet_paths:
        print(f"Sheet CSV: {sheet_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())