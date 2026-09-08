#!/usr/bin/env python3
"""Import bank CSV transactions into the existing Financements workbook."""

from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_WORKBOOK = "Financements.xlsx"
DEFAULT_RULES = "categories.json"

EXPENSE_LABEL_COL = 7  # G
EXPENSE_AMOUNT_COL = 8  # H
EXPENSE_INCLUDED_COL = 9  # I
EXPENSE_CATEGORY_COL = 11  # K

INCOME_LABEL_COL = 2  # B
INCOME_AMOUNT_COL = 3  # C

HEADER_ALIASES = {
    "date": ("date", "date operation", "date d operation", "date comptable"),
    "label": ("libelle", "operation", "description", "intitule", "designation"),
    "amount": ("montant", "amount", "valeur"),
    "debit": ("debit", "retrait", "depense"),
    "credit": ("credit", "versement", "recette"),
}


@dataclass(frozen=True)
class Transaction:
    date: date | None
    label: str
    amount: Decimal


def normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.strip().lower())
    ascii_value = "".join(char for char in decomposed if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", ascii_value).strip()


def parse_amount(value: Any) -> Decimal | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    cleaned = text.replace("\u00a0", " ").replace(" ", "")
    cleaned = cleaned.replace("EUR", "").replace("€", "")
    cleaned = cleaned.replace(",", ".")
    cleaned = re.sub(r"[^0-9.\-+]", "", cleaned)
    if not cleaned or cleaned in {"-", "+"}:
        return None
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def parse_date(value: Any) -> date | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def open_csv(path: Path) -> tuple[list[dict[str, str]], dict[str, str]]:
    for encoding in ("utf-8-sig", "cp1252", "iso-8859-1"):
        try:
            content = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Unable to decode {path}")

    sample = content[:4096]
    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    reader = csv.DictReader(content.splitlines(), dialect=dialect)
    if not reader.fieldnames:
        raise ValueError("CSV file has no header row")

    normalized_headers = {normalize(header): header for header in reader.fieldnames if header}
    return list(reader), normalized_headers


def find_header(headers: dict[str, str], logical_name: str) -> str | None:
    for alias in HEADER_ALIASES[logical_name]:
        normalized_alias = normalize(alias)
        if normalized_alias in headers:
            return headers[normalized_alias]
    return None


def read_transactions(path: Path) -> list[Transaction]:
    rows, headers = open_csv(path)
    date_header = find_header(headers, "date")
    label_header = find_header(headers, "label")
    amount_header = find_header(headers, "amount")
    debit_header = find_header(headers, "debit")
    credit_header = find_header(headers, "credit")

    if not label_header:
        raise ValueError("Unable to find a label column. Expected one of: libelle, operation, description")
    if not amount_header and not (debit_header or credit_header):
        raise ValueError("Unable to find an amount column. Expected montant, or debit/credit columns")

    transactions: list[Transaction] = []
    for row in rows:
        label = " ".join(str(row.get(label_header, "")).split())
        if not label:
            continue

        if amount_header:
            amount = parse_amount(row.get(amount_header))
        else:
            debit = parse_amount(row.get(debit_header)) if debit_header else None
            credit = parse_amount(row.get(credit_header)) if credit_header else None
            amount = (credit or Decimal("0")) - (debit or Decimal("0"))

        if amount is None or amount == 0:
            continue
        transactions.append(Transaction(parse_date(row.get(date_header)) if date_header else None, label, amount))
    return transactions


def load_rules(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise ValueError("categories.json must contain a JSON array")
    return data


def guess_category(label: str, rules: Iterable[dict[str, Any]], default: str) -> tuple[str, bool]:
    normalized_label = normalize(label)
    for rule in rules:
        patterns = rule.get("match", [])
        if isinstance(patterns, str):
            patterns = [patterns]
        for pattern in patterns:
            if normalize(str(pattern)) in normalized_label:
                return str(rule.get("category", default)), bool(rule.get("include", True))
    return default, True


def existing_values(ws: Any, min_row: int, max_row: int, label_col: int, amount_col: int) -> set[tuple[str, Decimal]]:
    values: set[tuple[str, Decimal]] = set()
    for row in range(min_row, max_row + 1):
        label = ws.cell(row=row, column=label_col).value
        amount = parse_amount(ws.cell(row=row, column=amount_col).value)
        if label and amount is not None:
            values.add((normalize(str(label)), amount.copy_abs()))
    return values


def first_empty_row(ws: Any, min_row: int, max_row: int, label_col: int) -> int | None:
    for row in range(min_row, max_row + 1):
        if ws.cell(row=row, column=label_col).value in (None, ""):
            return row
    return None


def following_empty_row(ws: Any, current_row: int, max_row: int, label_col: int) -> int | None:
    return first_empty_row(ws, current_row + 1, max_row, label_col)


def month_filter(transactions: list[Transaction], month: str | None) -> list[Transaction]:
    if not month:
        return transactions
    wanted = datetime.strptime(month, "%Y-%m").date()
    return [tx for tx in transactions if tx.date and tx.date.year == wanted.year and tx.date.month == wanted.month]


def import_transactions(args: argparse.Namespace) -> int:
    workbook_path = Path(args.workbook)
    csv_path = Path(args.csv_file)
    rules = load_rules(Path(args.rules))
    transactions = month_filter(read_transactions(csv_path), args.month)

    workbook = load_workbook(workbook_path)
    if args.sheet not in workbook.sheetnames:
        raise ValueError(f"Sheet '{args.sheet}' does not exist in {workbook_path}")
    ws = workbook[args.sheet]

    existing_expenses = existing_values(ws, args.expense_start, args.expense_end, EXPENSE_LABEL_COL, EXPENSE_AMOUNT_COL)
    existing_incomes = existing_values(ws, args.income_start, args.income_end, INCOME_LABEL_COL, INCOME_AMOUNT_COL)
    next_expense_row = first_empty_row(ws, args.expense_start, args.expense_end, EXPENSE_LABEL_COL)
    next_income_row = first_empty_row(ws, args.income_start, args.income_end, INCOME_LABEL_COL)

    imported = 0
    skipped_duplicates = 0
    skipped_no_space = 0
    preview: list[str] = []

    for tx in transactions:
        amount_abs = tx.amount.copy_abs()
        if tx.amount < 0:
            key = (normalize(tx.label), amount_abs)
            if key in existing_expenses:
                skipped_duplicates += 1
                continue
            if next_expense_row is None:
                skipped_no_space += 1
                continue
            target_row = next_expense_row
            category, include = guess_category(tx.label, rules, args.default_category)
            preview.append(f"depense row {target_row}: {tx.label} | {amount_abs} | {category}")
            if not args.dry_run:
                ws.cell(row=target_row, column=EXPENSE_LABEL_COL, value=tx.label)
                ws.cell(row=target_row, column=EXPENSE_AMOUNT_COL, value=float(amount_abs))
                ws.cell(row=target_row, column=EXPENSE_INCLUDED_COL, value="x" if include else None)
                ws.cell(row=target_row, column=EXPENSE_CATEGORY_COL, value=category)
            existing_expenses.add(key)
            next_expense_row = following_empty_row(ws, target_row, args.expense_end, EXPENSE_LABEL_COL)
            imported += 1
        else:
            key = (normalize(tx.label), amount_abs)
            if key in existing_incomes:
                skipped_duplicates += 1
                continue
            if next_income_row is None:
                skipped_no_space += 1
                continue
            target_row = next_income_row
            preview.append(f"recette row {target_row}: {tx.label} | {amount_abs}")
            if not args.dry_run:
                ws.cell(row=target_row, column=INCOME_LABEL_COL, value=tx.label)
                ws.cell(row=target_row, column=INCOME_AMOUNT_COL, value=float(amount_abs))
            existing_incomes.add(key)
            next_income_row = following_empty_row(ws, target_row, args.income_end, INCOME_LABEL_COL)
            imported += 1

    for line in preview[: args.preview_limit]:
        print(line)
    if len(preview) > args.preview_limit:
        print(f"... {len(preview) - args.preview_limit} more")

    print(f"Imported: {imported}")
    print(f"Skipped duplicates: {skipped_duplicates}")
    print(f"Skipped because target range is full: {skipped_no_space}")

    if not args.dry_run:
        workbook.save(workbook_path)
        print(f"Saved: {workbook_path}")
    else:
        print("Dry run only: workbook was not modified")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Import bank CSV transactions into Financements.xlsx")
    parser.add_argument("csv_file", help="Bank CSV export to import")
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK, help="Workbook path")
    parser.add_argument("--sheet", required=True, help="Monthly sheet name, for example 'Septembre 2026'")
    parser.add_argument("--month", help="Optional CSV date filter in YYYY-MM format")
    parser.add_argument("--rules", default=DEFAULT_RULES, help="Categorization rules JSON file")
    parser.add_argument("--default-category", default="Divers", help="Category used when no rule matches")
    parser.add_argument("--expense-start", type=int, default=44, help="First row for variable expenses")
    parser.add_argument("--expense-end", type=int, default=247, help="Last row for variable expenses")
    parser.add_argument("--income-start", type=int, default=8, help="First row for additional incomes")
    parser.add_argument("--income-end", type=int, default=35, help="Last row for additional incomes")
    parser.add_argument("--preview-limit", type=int, default=25, help="Maximum preview lines to print")
    parser.add_argument("--dry-run", action="store_true", help="Preview changes without saving the workbook")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    return import_transactions(args)


if __name__ == "__main__":
    raise SystemExit(main())