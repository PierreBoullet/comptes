#!/usr/bin/env python3
"""Create a dashboard-friendly CSV from Financements.xlsx."""

from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MONTHS = {
    "janvier": 1,
    "fevrier": 2,
    "février": 2,
    "mars": 3,
    "avril": 4,
    "mai": 5,
    "juin": 6,
    "juillet": 7,
    "aout": 8,
    "août": 8,
    "septembre": 9,
    "octobre": 10,
    "novembre": 11,
    "decembre": 12,
    "décembre": 12,
}


def normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.strip().lower())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def month_from_sheet(sheet_name: str) -> date | None:
    match = re.match(r"^([A-Za-zéûôîàèùç]+)\s+(\d{4})$", sheet_name.strip())
    if not match:
        return None
    month = MONTHS.get(normalize(match.group(1)))
    if not month:
        return None
    return date(int(match.group(2)), month, 1)


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\u00a0", " ").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def rows_from_workbook(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        operation_date = month_from_sheet(worksheet.title)
        if operation_date is None:
            continue

        iso_date = operation_date.isoformat()

        for row_index, row in enumerate(worksheet.iter_rows(min_row=2, max_row=35, min_col=2, max_col=3, values_only=True), start=2):
            label, raw_amount = row
            amount = number(raw_amount)
            if label and amount:
                rows.append({
                    "Date": iso_date,
                    "Libellé": str(label).strip(),
                    "Montant": round(abs(amount), 2),
                    "Catégorie": "Recettes",
                    "Type": "recette",
                    "Source": "Financements.xlsx",
                    "Feuille": worksheet.title,
                    "Ligne": row_index,
                })

        for row_index, row in enumerate(worksheet.iter_rows(min_row=3, max_row=38, min_col=5, max_col=11, values_only=True), start=3):
            label = row[0]
            amount = number(row[3])
            category = row[6] or "Dépenses prévues"
            if label and amount:
                rows.append({
                    "Date": iso_date,
                    "Libellé": str(label).strip(),
                    "Montant": round(-abs(amount), 2),
                    "Catégorie": str(category).strip(),
                    "Type": "depense_prevue",
                    "Source": "Financements.xlsx",
                    "Feuille": worksheet.title,
                    "Ligne": row_index,
                })

        for row_index, row in enumerate(worksheet.iter_rows(min_row=44, max_row=247, min_col=7, max_col=11, values_only=True), start=44):
            label = row[0]
            amount = number(row[3])
            category = row[4] or "Divers"
            if label and amount:
                rows.append({
                    "Date": iso_date,
                    "Libellé": str(label).strip(),
                    "Montant": round(-abs(amount), 2),
                    "Catégorie": str(category).strip(),
                    "Type": "depense_non_prevue",
                    "Source": "Financements.xlsx",
                    "Feuille": worksheet.title,
                    "Ligne": row_index,
                })

    return rows


def write_csv(rows: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = ["Date", "Libellé", "Montant", "Catégorie", "Type", "Source", "Feuille", "Ligne"]
    with output_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export Financements.xlsx transactions to data CSV")
    parser.add_argument("--workbook", default="Financements.xlsx", help="Workbook to export")
    parser.add_argument("--output", default="data/financements_transactions.csv", help="CSV output path")
    args = parser.parse_args()

    rows = rows_from_workbook(Path(args.workbook))
    write_csv(rows, Path(args.output))
    print(f"Exported {len(rows)} transactions to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())